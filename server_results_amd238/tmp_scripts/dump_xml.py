"""对比模板与生成文件里同一形状的 XML，用于定位渲染差异。"""

import sys

from pptx import Presentation
from lxml import etree

path, idx, name = sys.argv[1], int(sys.argv[2]), sys.argv[3]
prs = Presentation(path)
for shp in prs.slides[idx].shapes:
    if shp.name == name:
        print(etree.tostring(shp._element, pretty_print=True).decode("utf-8"))
        break
else:
    print("未找到", name, [s.name for s in prs.slides[idx].shapes])