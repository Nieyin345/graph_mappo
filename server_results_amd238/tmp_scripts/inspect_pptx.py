"""导出 极简黑白.pptx 每页的形状 / 几何 / 文本 / 字体，用于决定构建方式。"""

from pptx import Presentation


def cm(v):
    return None if v is None else round(v / 360000, 2)


def geo(sh):
    return f"({cm(sh.left)}, {cm(sh.top)}, {cm(sh.width)} x {cm(sh.height)})"


def fontinfo(sh):
    out = []
    if not sh.has_text_frame:
        return ""
    for par in sh.text_frame.paragraphs:
        for r in par.runs:
            c = None
            try:
                c = r.font.color.rgb
            except Exception:
                c = "scheme"
            out.append((r.text[:24], r.font.size.pt if r.font.size else None,
                        r.font.name, r.font.bold, str(c)))
    return out


def walk(shapes, indent="  "):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            print(f"{indent}[GROUP] {sh.name} {geo(sh)}")
            walk(sh.shapes, indent + "    ")
        else:
            txt = ""
            if sh.has_text_frame and sh.text_frame.text.strip():
                txt = repr(sh.text_frame.text[:60].replace("\n", " / "))
            print(f"{indent}{sh.shape_type} {sh.name} {geo(sh)} {txt}")
            fi = fontinfo(sh)
            if fi:
                print(f"{indent}    fonts: {fi[:4]}")


prs = Presentation("极简黑白.pptx")
print("SLIDE SIZE:", cm(prs.slide_width), "x", cm(prs.slide_height))
for i, slide in enumerate(prs.slides):
    print(f"\n===== SLIDE {i} | layout={slide.slide_layout.name} =====")
    walk(slide.shapes)